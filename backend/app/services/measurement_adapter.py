import io
import json
from datetime import datetime

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import MeasurementFieldMapping, MeasurementRecord, Monitor, Network
from app.models.measurement import SIGNAL_FIELDS
from app.schemas.measurement import MeasurementImportResult, ResolveUnmatchedResult, UnmatchedNameRow
from app.services.monitor_matcher import match_monitor_name

# Section 14.1 default mapping for the current production Excel export.
# Configurable per network via the field-mapping API -- never hardcoded into
# the import pipeline itself.
DEFAULT_FIELD_MAPPING: dict[str, str] = {
    "monitor_name": "监测点名称1",
    "device_type": "设备类型",
    "terminal_status": "终端状态",
    "line_status": "线路状态",
    "warning_status": "预警状态",
    "Va": "A相电压(kV)",
    "Vb": "B相电压(kV)",
    "Vc": "C相电压(kV)",
    "Ia": "A相电流(A)",
    "Ib": "B相电流(A)",
    "Ic": "C相电流(A)",
    "phase_Va": "A相电压相位",
    "phase_Vb": "B相电压相位",
    "phase_Vc": "C相电压相位",
    "phase_Ia": "A相电流相位",
    "phase_Ib": "B相电流相位",
    "phase_Ic": "C相电流相位",
    "timestamp": "量测时间",
}

TARGET_FIELDS = list(DEFAULT_FIELD_MAPPING.keys())

INVALID_NUMERIC_SENTINELS = {"-9.999", "-9999", "---", "--", "-", ""}
INVALID_NUMERIC_VALUES = {-9.999, -9999.0}


class MeasurementImportError(ValueError):
    pass


def get_field_mapping(db: Session, network_id: str) -> dict[str, str]:
    rows = db.query(MeasurementFieldMapping).filter(MeasurementFieldMapping.network_id == network_id).all()
    if not rows:
        return dict(DEFAULT_FIELD_MAPPING)
    configured = {r.target_field: r.source_column for r in rows}
    return {**DEFAULT_FIELD_MAPPING, **configured}


def set_field_mapping(db: Session, network_id: str, mapping: dict[str, str]) -> dict[str, str]:
    unknown = set(mapping) - set(TARGET_FIELDS)
    if unknown:
        raise MeasurementImportError(f"未知的标准字段：{', '.join(sorted(unknown))}")

    db.query(MeasurementFieldMapping).filter(MeasurementFieldMapping.network_id == network_id).delete()
    for target_field, source_column in mapping.items():
        db.add(MeasurementFieldMapping(network_id=network_id, target_field=target_field, source_column=source_column))
    db.commit()
    return get_field_mapping(db, network_id)


def clean_numeric(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return None if float(value) in INVALID_NUMERIC_VALUES else float(value)
    text = str(value).strip()
    if text in INVALID_NUMERIC_SENTINELS:
        return None
    try:
        parsed = float(text)
    except ValueError:
        return None
    return None if parsed in INVALID_NUMERIC_VALUES else parsed


def _clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_timestamp(value: object) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def import_measurements(
    db: Session,
    network: Network,
    file_bytes: bytes,
    sheet_name: str | None = None,
) -> MeasurementImportResult:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]

    header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
    header = [str(c.value).strip() if c.value is not None else "" for c in header_row]
    column_index = {name: idx for idx, name in enumerate(header) if name}

    mapping = get_field_mapping(db, network.network_id)
    if "monitor_name" not in mapping or mapping["monitor_name"] not in column_index:
        raise MeasurementImportError(
            f"未找到监测点名称列 '{mapping.get('monitor_name')}'，请检查字段映射配置"
        )

    def get(values: list, target_field: str) -> object | None:
        source_column = mapping.get(target_field)
        if source_column is None:
            return None
        idx = column_index.get(source_column)
        if idx is None or idx >= len(values):
            return None
        return values[idx]

    matched_count = 0
    unmatched_count = 0
    unmatched_names: set[str] = set()
    total_rows = 0

    for row in sheet.iter_rows(min_row=2):
        values = [cell.value for cell in row]
        if all(v is None for v in values):
            continue

        monitor_name_raw = _clean_text(get(values, "monitor_name"))
        if monitor_name_raw is None:
            continue

        total_rows += 1
        monitor_id = match_monitor_name(db, network.network_id, monitor_name_raw)
        match_status = "matched" if monitor_id else "unmatched"
        if monitor_id:
            matched_count += 1
        else:
            unmatched_count += 1
            unmatched_names.add(monitor_name_raw)

        record = MeasurementRecord(
            network_id=network.network_id,
            monitor_name_raw=monitor_name_raw,
            monitor_id=monitor_id,
            match_status=match_status,
            device_type=_clean_text(get(values, "device_type")),
            terminal_status=_clean_text(get(values, "terminal_status")),
            line_status=_clean_text(get(values, "line_status")),
            warning_status=_clean_text(get(values, "warning_status")),
            timestamp=_parse_timestamp(get(values, "timestamp")),
            raw_json=json.dumps(dict(zip(header, values, strict=False)), default=str, ensure_ascii=False),
            **{signal: clean_numeric(get(values, signal)) for signal in SIGNAL_FIELDS},
        )
        db.add(record)

    db.commit()

    return MeasurementImportResult(
        network_id=network.network_id,
        total_rows=total_rows,
        matched=matched_count,
        unmatched=unmatched_count,
        unmatched_names=sorted(unmatched_names),
    )


def list_unmatched_names(db: Session, network_id: str) -> list[UnmatchedNameRow]:
    rows = (
        db.query(MeasurementRecord.monitor_name_raw, MeasurementRecord.id)
        .filter(MeasurementRecord.network_id == network_id, MeasurementRecord.match_status == "unmatched")
        .all()
    )
    counts: dict[str, int] = {}
    for name, _id in rows:
        counts[name] = counts.get(name, 0) + 1
    return [UnmatchedNameRow(monitor_name_raw=name, count=count) for name, count in sorted(counts.items())]


def resolve_unmatched(
    db: Session, network: Network, monitor_name_raw: str, monitor_id: str | None, ignore: bool
) -> ResolveUnmatchedResult:
    if ignore:
        updated = (
            db.query(MeasurementRecord)
            .filter(
                MeasurementRecord.network_id == network.network_id,
                MeasurementRecord.monitor_name_raw == monitor_name_raw,
                MeasurementRecord.match_status == "unmatched",
            )
            .update({"match_status": "ignored"})
        )
        db.commit()
        return ResolveUnmatchedResult(monitor_name_raw=monitor_name_raw, updated_records=updated, status="ignored")

    if not monitor_id:
        raise MeasurementImportError("请指定 monitor_id，或设置 ignore=true")

    monitor = db.get(Monitor, monitor_id)
    if monitor is None or monitor.network_id != network.network_id:
        raise MeasurementImportError("指定的 Monitor 不存在于该 Network 中")

    if monitor_name_raw not in monitor.aliases and monitor_name_raw != monitor.canonical_name:
        monitor.aliases = [*monitor.aliases, monitor_name_raw]

    updated = (
        db.query(MeasurementRecord)
        .filter(
            MeasurementRecord.network_id == network.network_id,
            MeasurementRecord.monitor_name_raw == monitor_name_raw,
            MeasurementRecord.match_status == "unmatched",
        )
        .update({"match_status": "matched", "monitor_id": monitor_id})
    )
    db.commit()
    return ResolveUnmatchedResult(monitor_name_raw=monitor_name_raw, updated_records=updated, status="matched")
