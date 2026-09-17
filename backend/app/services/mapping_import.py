import io
from dataclasses import dataclass

import networkx as nx
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.db.id_generator import next_id
from app.models import Edge, Network, Node
from app.schemas.mapping import MappingImportResult, ValidationIssue
from app.services.text_normalize import normalize_text

VALID_STATUSES = {"closed", "open"}


@dataclass
class RawEdgeRow:
    row_number: int  # 1-based Excel row, header is row 1
    edge_key: str | None
    node_1: str | None
    node_2: str | None
    status: str | None
    remark: str | None
    node_1_trimmed: bool = False
    node_2_trimmed: bool = False


class MappingImportError(ValueError):
    pass


def _clean(value: object) -> tuple[str | None, bool]:
    """Returns (trimmed value or None, whether trimming changed a non-empty string)."""
    if value is None:
        return None, False
    raw = str(value)
    trimmed = raw.strip()
    if not trimmed:
        return None, False
    return trimmed, trimmed != raw


def parse_edge_rows(file_bytes: bytes, sheet_name: str = "Edges") -> list[RawEdgeRow]:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise MappingImportError(f"未找到工作表 '{sheet_name}'，可用工作表：{workbook.sheetnames}")
    sheet = workbook[sheet_name]

    header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
    header = [str(c.value).strip() if c.value is not None else "" for c in header_row]
    column_index = {name: idx for idx, name in enumerate(header) if name}

    for required in ("node_1", "node_2"):
        if required not in column_index:
            raise MappingImportError(f"工作表 '{sheet_name}' 缺少必需列 '{required}'")

    def get(values: list, col: str) -> object | None:
        idx = column_index.get(col)
        if idx is None or idx >= len(values):
            return None
        return values[idx]

    rows: list[RawEdgeRow] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = [cell.value for cell in row]

        node_1, node_1_trimmed = _clean(get(values, "node_1"))
        node_2, node_2_trimmed = _clean(get(values, "node_2"))
        edge_key, _ = _clean(get(values, "edge_key"))
        status, _ = _clean(get(values, "status"))
        remark, _ = _clean(get(values, "remark"))

        if not any((node_1, node_2, edge_key, status, remark)):
            continue

        rows.append(
            RawEdgeRow(
                row_number=row_number,
                edge_key=edge_key,
                node_1=node_1,
                node_2=node_2,
                status=status.lower() if status else status,
                remark=remark,
                node_1_trimmed=node_1_trimmed,
                node_2_trimmed=node_2_trimmed,
            )
        )

    return rows


def validate_rows(rows: list[RawEdgeRow]) -> tuple[list[ValidationIssue], list[ValidationIssue]]:
    errors: list[ValidationIssue] = []
    warnings: list[ValidationIssue] = []

    missing_1 = [r.row_number for r in rows if r.node_1 is None]
    if missing_1:
        errors.append(ValidationIssue(level="ERROR", code="missing_node_1", message="node_1 为空", rows=missing_1))

    missing_2 = [r.row_number for r in rows if r.node_2 is None]
    if missing_2:
        errors.append(ValidationIssue(level="ERROR", code="missing_node_2", message="node_2 为空", rows=missing_2))

    self_loops = [r.row_number for r in rows if r.node_1 and r.node_2 and r.node_1 == r.node_2]
    if self_loops:
        errors.append(
            ValidationIssue(level="ERROR", code="self_loop", message="node_1 == node_2（自环）", rows=self_loops)
        )

    pair_rows: dict[frozenset[str], list[int]] = {}
    for r in rows:
        if r.node_1 and r.node_2 and r.node_1 != r.node_2:
            pair_rows.setdefault(frozenset((r.node_1, r.node_2)), []).append(r.row_number)
    duplicate_pair_rows = sorted(row for group in pair_rows.values() if len(group) > 1 for row in group)
    if duplicate_pair_rows:
        errors.append(
            ValidationIssue(
                level="ERROR",
                code="duplicate_edge",
                message="存在完全重复的无向 Edge（含 A-B 与 B-A 反向重复）",
                rows=duplicate_pair_rows,
            )
        )

    key_rows: dict[str, list[int]] = {}
    for r in rows:
        if r.edge_key:
            key_rows.setdefault(r.edge_key, []).append(r.row_number)
    duplicate_key_rows = sorted(row for group in key_rows.values() if len(group) > 1 for row in group)
    if duplicate_key_rows:
        errors.append(
            ValidationIssue(
                level="ERROR", code="duplicate_edge_key", message="edge_key 重复（非空）", rows=duplicate_key_rows
            )
        )

    invalid_status = [r.row_number for r in rows if r.status is not None and r.status not in VALID_STATUSES]
    if invalid_status:
        errors.append(
            ValidationIssue(
                level="ERROR", code="invalid_status", message="status 不属于 closed/open/空", rows=invalid_status
            )
        )

    if errors:
        return errors, warnings

    graph = nx.Graph()
    for r in rows:
        graph.add_edge(r.node_1, r.node_2)

    components = nx.number_connected_components(graph) if graph.number_of_nodes() else 0
    if components > 1:
        warnings.append(
            ValidationIssue(level="WARNING", code="multiple_components", message=f"存在 {components} 个连通分量")
        )

    cyclomatic = graph.number_of_edges() - graph.number_of_nodes() + components
    if cyclomatic > 0:
        warnings.append(
            ValidationIssue(level="WARNING", code="cycle_detected", message=f"图中存在 {cyclomatic} 个独立环路")
        )

    warnings.append(ValidationIssue(level="WARNING", code="source_not_set", message="尚未选择 Source"))

    open_rows = [r.row_number for r in rows if r.status == "open"]
    if open_rows:
        warnings.append(ValidationIssue(level="WARNING", code="open_edge", message="存在 open Edge", rows=open_rows))

    whitespace_rows = sorted(
        r.row_number for r in rows if r.node_1_trimmed or r.node_2_trimmed
    )
    if whitespace_rows:
        warnings.append(
            ValidationIssue(
                level="WARNING",
                code="node_key_whitespace",
                message="Node Key 包含前后空格，已自动去除",
                rows=whitespace_rows,
            )
        )

    normalized_groups: dict[str, set[str]] = {}
    for key in {*(r.node_1 for r in rows if r.node_1), *(r.node_2 for r in rows if r.node_2)}:
        normalized_groups.setdefault(normalize_text(key), set()).add(key)
    ambiguous = [group for group in normalized_groups.values() if len(group) > 1]
    if ambiguous:
        sample = "; ".join("/".join(sorted(group)) for group in ambiguous)
        warnings.append(
            ValidationIssue(
                level="WARNING",
                code="node_key_suspected_duplicate",
                message=f"Node Key 疑似重复格式，请人工确认是否为同一节点：{sample}",
            )
        )

    return errors, warnings


def import_mapping(
    db: Session,
    network: Network,
    file_bytes: bytes,
    sheet_name: str = "Edges",
) -> MappingImportResult:
    existing_node_count = db.query(Node).filter(Node.network_id == network.network_id).count()
    if existing_node_count > 0:
        raise MappingImportError("该 Network 已存在拓扑数据，请创建新 Network 后重新导入")

    rows = parse_edge_rows(file_bytes, sheet_name=sheet_name)
    errors, warnings = validate_rows(rows)

    if errors:
        return MappingImportResult(
            imported=False,
            network_id=network.network_id,
            nodes=0,
            edges=0,
            components=0,
            cycles=0,
            open_edges=0,
            errors=errors,
            warnings=warnings,
        )

    node_ids: dict[str, str] = {}
    for r in rows:
        for key in (r.node_1, r.node_2):
            if key not in node_ids:
                node = Node(node_id=next_id(db, "N"), network_id=network.network_id, node_key=key)
                db.add(node)
                db.flush()
                node_ids[key] = node.node_id

    open_edge_count = 0
    for r in rows:
        status = r.status or "closed"
        if status == "open":
            open_edge_count += 1
        edge = Edge(
            edge_id=next_id(db, "E"),
            network_id=network.network_id,
            edge_key=r.edge_key,
            node_a_id=node_ids[r.node_1],
            node_b_id=node_ids[r.node_2],
            status=status,
            remark=r.remark,
        )
        db.add(edge)

    db.commit()

    graph = nx.Graph()
    graph.add_nodes_from(node_ids.values())
    for r in rows:
        graph.add_edge(node_ids[r.node_1], node_ids[r.node_2])
    components = nx.number_connected_components(graph)
    cyclomatic = graph.number_of_edges() - graph.number_of_nodes() + components

    return MappingImportResult(
        imported=True,
        network_id=network.network_id,
        nodes=len(node_ids),
        edges=len(rows),
        components=components,
        cycles=cyclomatic,
        open_edges=open_edge_count,
        errors=[],
        warnings=warnings,
    )
