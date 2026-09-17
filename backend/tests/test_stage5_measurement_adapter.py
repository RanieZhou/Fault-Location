import io

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models import MeasurementRecord
from app.services.measurement_adapter import DEFAULT_FIELD_MAPPING, clean_numeric

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

PRODUCTION_HEADER = list(DEFAULT_FIELD_MAPPING.values())


def _make_production_workbook(rows: list[dict]) -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"
    sheet.append(PRODUCTION_HEADER)
    for row in rows:
        sheet.append([row.get(col, None) for col in PRODUCTION_HEADER])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _row(monitor_name, **overrides) -> dict:
    base = {
        "监测点名称1": monitor_name,
        "设备类型": "配电线路",
        "终端状态": "正常",
        "线路状态": "正常",
        "预警状态": "正常",
        "A相电压(kV)": 10.01,
        "B相电压(kV)": 10.02,
        "C相电压(kV)": 10.03,
        "A相电流(A)": 5.1,
        "B相电流(A)": 5.2,
        "C相电流(A)": 5.3,
        "A相电压相位": 0.0,
        "B相电压相位": -120.0,
        "C相电压相位": 120.0,
        "A相电流相位": 1.0,
        "B相电流相位": -119.0,
        "C相电流相位": 121.0,
        "量测时间": "2026-01-01 00:00:00",
    }
    base.update(overrides)
    return base


def _create_network(client: TestClient) -> str:
    return client.post("/api/networks", json={"name": "Test Feeder"}).json()["network_id"]


def test_default_field_mapping_returned_when_not_configured(client: TestClient) -> None:
    network_id = _create_network(client)
    response = client.get(f"/api/networks/{network_id}/measurements/mapping")
    assert response.status_code == 200
    assert response.json()["mapping"] == DEFAULT_FIELD_MAPPING


def test_field_mapping_override_merges_with_defaults(client: TestClient) -> None:
    network_id = _create_network(client)
    response = client.post(
        f"/api/networks/{network_id}/measurements/mapping",
        json={"mapping": {"monitor_name": "监测点名称"}},
    )
    assert response.status_code == 200
    mapping = response.json()["mapping"]
    assert mapping["monitor_name"] == "监测点名称"
    assert mapping["Va"] == DEFAULT_FIELD_MAPPING["Va"]


def test_clean_numeric_rejects_known_sentinels() -> None:
    assert clean_numeric("-9.999") is None
    assert clean_numeric(-9.999) is None
    assert clean_numeric("-9999") is None
    assert clean_numeric(-9999) is None
    assert clean_numeric("---") is None
    assert clean_numeric("") is None
    assert clean_numeric(None) is None
    assert clean_numeric("10.5") == 10.5
    assert clean_numeric(10.5) == 10.5


def test_import_matches_canonical_and_alias_and_cleans_invalid_values(
    client: TestClient, db_session: Session
) -> None:
    network_id = _create_network(client)
    file_bytes = io.BytesIO()
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Edges"
    sheet.append(["edge_key", "node_1", "node_2", "status", "remark"])
    sheet.append(["L001", "#1", "#2", "closed", ""])
    wb.save(file_bytes)
    client.post(
        f"/api/networks/{network_id}/mapping/import",
        files={"file": ("mapping.xlsx", file_bytes.getvalue(), XLSX_CONTENT_TYPE)},
    )
    node_id = client.get(f"/api/networks/{network_id}/graph").json()["nodes"][0]["node_id"]
    client.post(
        f"/api/networks/{network_id}/monitors",
        json={"node_id": node_id, "canonical_name": "10kV火龙线#67.7支", "aliases": ["火龙线67.7支"]},
    )

    rows = [
        _row("10kV火龙线#67.7支"),  # exact canonical match
        _row(" 火龙线67.7支 "),  # alias match, tolerant of surrounding whitespace
        _row("未知监测点", **{"A相电压(kV)": "-9.999", "B相电流(A)": -9999, "C相电流(A)": "---"}),
    ]
    measurement_bytes = _make_production_workbook(rows)

    response = client.post(
        f"/api/networks/{network_id}/measurements/import",
        files={"file": ("prod.xlsx", measurement_bytes, XLSX_CONTENT_TYPE)},
    )
    assert response.status_code == 200
    result = response.json()
    assert result["total_rows"] == 3
    assert result["matched"] == 2
    assert result["unmatched"] == 1
    assert result["unmatched_names"] == ["未知监测点"]

    unmatched_record = (
        db_session.query(MeasurementRecord)
        .filter(MeasurementRecord.network_id == network_id, MeasurementRecord.monitor_name_raw == "未知监测点")
        .one()
    )
    assert unmatched_record.Va is None
    assert unmatched_record.Ib is None
    assert unmatched_record.Ic is None
    assert unmatched_record.match_status == "unmatched"

    matched_record = (
        db_session.query(MeasurementRecord)
        .filter(MeasurementRecord.network_id == network_id, MeasurementRecord.monitor_name_raw == "10kV火龙线#67.7支")
        .one()
    )
    assert matched_record.match_status == "matched"
    assert matched_record.Va == 10.01


def test_unmatched_review_resolve_binds_alias_and_updates_records(client: TestClient) -> None:
    network_id = _create_network(client)
    file_bytes = io.BytesIO()
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Edges"
    sheet.append(["edge_key", "node_1", "node_2", "status", "remark"])
    sheet.append(["L001", "#1", "#2", "closed", ""])
    wb.save(file_bytes)
    client.post(
        f"/api/networks/{network_id}/mapping/import",
        files={"file": ("mapping.xlsx", file_bytes.getvalue(), XLSX_CONTENT_TYPE)},
    )
    node_id = client.get(f"/api/networks/{network_id}/graph").json()["nodes"][0]["node_id"]
    monitor = client.post(
        f"/api/networks/{network_id}/monitors", json={"node_id": node_id, "canonical_name": "M-A"}
    ).json()

    measurement_bytes = _make_production_workbook([_row("M-A-typo")])
    client.post(
        f"/api/networks/{network_id}/measurements/import",
        files={"file": ("prod.xlsx", measurement_bytes, XLSX_CONTENT_TYPE)},
    )

    unmatched = client.get(f"/api/networks/{network_id}/measurements/unmatched").json()
    assert unmatched == [{"monitor_name_raw": "M-A-typo", "count": 1}]

    resolve_resp = client.post(
        f"/api/networks/{network_id}/measurements/resolve",
        json={"monitor_name_raw": "M-A-typo", "monitor_id": monitor["monitor_id"]},
    )
    assert resolve_resp.status_code == 200
    assert resolve_resp.json()["updated_records"] == 1
    assert resolve_resp.json()["status"] == "matched"

    assert client.get(f"/api/networks/{network_id}/measurements/unmatched").json() == []

    monitors = client.get(f"/api/networks/{network_id}/monitors").json()
    assert "M-A-typo" in monitors[0]["aliases"]


def test_resolve_with_ignore_marks_ignored_without_binding(client: TestClient) -> None:
    network_id = _create_network(client)
    file_bytes = io.BytesIO()
    wb = Workbook()
    sheet = wb.active
    sheet.append(["edge_key", "node_1", "node_2", "status", "remark"])
    sheet.append(["L001", "#1", "#2", "closed", ""])
    wb.save(file_bytes)
    client.post(
        f"/api/networks/{network_id}/mapping/import",
        files={"file": ("mapping.xlsx", file_bytes.getvalue(), XLSX_CONTENT_TYPE)},
    )

    measurement_bytes = _make_production_workbook([_row("废弃设备")])
    client.post(
        f"/api/networks/{network_id}/measurements/import",
        files={"file": ("prod.xlsx", measurement_bytes, XLSX_CONTENT_TYPE)},
    )

    resolve_resp = client.post(
        f"/api/networks/{network_id}/measurements/resolve",
        json={"monitor_name_raw": "废弃设备", "ignore": True},
    )
    assert resolve_resp.status_code == 200
    assert resolve_resp.json()["status"] == "ignored"
    assert client.get(f"/api/networks/{network_id}/measurements/unmatched").json() == []
