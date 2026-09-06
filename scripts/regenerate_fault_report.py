# -*- coding: utf-8 -*-
"""
用新版 core.fault_locator（Bug3滑动窗口 + 置信度聚合修复）重新生成
output/fault_location_report.csv。

与旧版 scripts/locate_fault.py 的输出差异：
  - 每个事件一行（旧版按"前沿点"展开，双前沿事件会拆成多行）；
    多个前沿/候选区段用 " / " 连接，和 web/static/js/data.js 里手工整理的格式一致。
  - 新增 confidence 列（high/medium/low），直接来自 core.fault_locator 的
    结构化计算，不再需要 core/data_loader.py 用正则/子串猜测 note 文本推断置信度。
运行前会自动备份旧文件为 fault_location_report.csv.bak（若尚不存在）。
"""
import sys, os, csv
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import openpyxl
from common import parse_monitor_name
from build_topology import to_dt, IN_SCOPE_LINES, DATA_DIR

from core.fault_locator import group_alarms_into_events, locate_fault
from core.models import FaultLocateRequest


def load_raw_alarms():
    wb = openpyxl.load_workbook(os.path.join(DATA_DIR, '异常数据.xlsx'), read_only=True, data_only=True)
    ws = wb['Sheet1']
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()

    alarms = []
    for r in rows:
        info = parse_monitor_name(r['监测点名称1'])
        if info is None or info['line'] not in IN_SCOPE_LINES:
            continue
        alarms.append({
            'line': info['line'],
            'pole': info['pole'],
            'time': to_dt(r['量测时间']),
            'fault_type': r['线路状态'],
        })
    return alarms


def main():
    out_path = ROOT / "output" / "fault_location_report.csv"
    bak_path = ROOT / "output" / "fault_location_report.csv.bak"
    if out_path.exists() and not bak_path.exists():
        bak_path.write_bytes(out_path.read_bytes())
        print(f"已备份旧文件到 {bak_path}")

    alarms = load_raw_alarms()
    events = group_alarms_into_events(alarms)
    events.sort(key=lambda ev: (ev[0]['line'], ev[0]['time']))

    rows = []
    for i, ev in enumerate(events, 1):
        line = ev[0]['line']
        time_str = str(ev[0]['time'])
        alarmed_poles = sorted({p['pole'] for p in ev})
        fault_types = sorted({p['fault_type'] for p in ev})

        req = FaultLocateRequest(line=line, alarm_points=alarmed_poles, fault_type=None)
        result = locate_fault(req)

        cand_str = " / ".join(
            f"{s.to_pole}" if s.to_pole != "(末端)" else "(无)"
            for s in result.candidate_sections
        ) or "(无)"

        rows.append({
            'event_id': i,
            'line': line,
            'time': time_str,
            'fault_types': '; '.join(fault_types),
            'alarmed_points': ', '.join(alarmed_poles),
            'frontier': ' / '.join(result.frontier_points),
            'candidate_sections': cand_str,
            'confidence': result.confidence,
            'note': result.note,
        })

    with open(out_path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=[
            'event_id', 'line', 'time', 'fault_types', 'alarmed_points',
            'frontier', 'candidate_sections', 'confidence', 'note',
        ])
        w.writeheader()
        w.writerows(rows)

    print(f"已重新生成 {out_path} ({len(rows)} 行，此前为按前沿点展开的16行)")


if __name__ == '__main__':
    main()
