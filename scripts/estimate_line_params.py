# -*- coding: utf-8 -*-
"""
给 edges.csv 里每条区段估算长度和阻抗参数(没有逐杆塔GIS数据,只能近似):

长度分配:不是简单总长度/区段数平均分,而是按"编号差"加权——同一条支线上相邻
两个监测点(如 27.3.84 -> 27.3.88)编号差是4,大概率比编号差是1的区段(如
81 -> 82)更长,用这个差值做权重去分总长度,比纯平均更合理一点,但仍然是
粗估,不是精确值。跨层级分支的边(比如27.3.88 -> 27.3.88.2.2.1,前面新增了
分支层级)没有可比的编号差,统一给一个较小的默认权重(算作短距离引下线)。

导线参数:2026-09-03根据南网10kV架空线路文献典型值更新(用户提供),正序
R1=0.125~0.17 Ω/km、L1=1.21~1.3 mH/km、C1=0.0096~0.012 uF/km,零序
R0=0.23~0.275 Ω/km、L0=3.7~5.478 mH/km、C0=0.0054~0.009 uF/km。主干取
区间偏粗导线一侧(R1较小),支线取偏细一侧(R1较大),零序不做主干/支线区分
(没有更细的数据支撑),统一取区间中值。之前用的LGJ-95/50猜测值已作废。
"""
import sys, os, re, csv
sys.path.insert(0, os.path.dirname(__file__))
import fitz
from collections import defaultdict
from common import path_sort_key, LINE_CODE
from build_topology import build_tree, SOURCE, IN_SCOPE_LINES, OUT_DIR

DIAGRAM_DIR = r'D:\Desktop\故障定位\data\单线图'
DIAGRAM_FILE = {'松坪线': '10kV松坪线.pdf', '火龙线': '10kV火龙线.pdf'}

F = 50  # Hz
# 正序(用户提供的南网10kV架空线文献区间, 干线/支线各取区间一侧)
CONDUCTOR_PARAMS = {
    '干线(正序R1=0.13欧姆/km)': {'r1_per_km': 0.13, 'l1_mH_per_km': 1.22},
    '支线(正序R1=0.16欧姆/km)': {'r1_per_km': 0.16, 'l1_mH_per_km': 1.28},
}
# 零序(不分干线/支线,取用户提供区间的中值): R0=0.23~0.275, L0=3.7~5.478
ZERO_SEQ = {'r0_per_km': 0.2525, 'l0_mH_per_km': 4.589}
# 正序/零序电容(uF/km),用户提供区间中值: C1=0.0096~0.012, C0=0.0054~0.009
CAP = {'c1_uF_per_km': 0.0108, 'c0_uF_per_km': 0.0072}
DEFAULT_BRANCH_WEIGHT = 2.0  # 跨层级分支边没有可比编号差时的默认权重(相当于"编号差2"的短距离)


def get_total_length_km(line):
    path = os.path.join(DIAGRAM_DIR, DIAGRAM_FILE[line])
    doc = fitz.open(path)
    text = doc[0].get_text()
    doc.close()
    m = re.search(r'总长度[:：]\s*([\d.]+)\s*km', text)
    if not m:
        raise RuntimeError(f'{line} 图纸里没找到总长度标注')
    return float(m.group(1))


def seg_numeric(seg):
    if '+' in seg:
        base, extra = seg.split('+', 1)
        return float(base) + float(extra) * 0.1
    return float(seg)


def edge_weight(parent, child):
    if parent == SOURCE:
        return DEFAULT_BRANCH_WEIGHT
    if parent[:-1] == child[:-1] and len(parent) == len(child):
        # 同一层级(同一支线)相邻,用编号差做权重
        gap = abs(seg_numeric(child[-1]) - seg_numeric(parent[-1]))
        return max(gap, 0.5)
    return DEFAULT_BRANCH_WEIGHT


def main():
    out_rows = []
    for line in IN_SCOPE_LINES:
        total_km = get_total_length_km(line)

        # 复用 build_topology 的方式重新读一遍该线路的 path 集合
        import openpyxl
        from common import parse_monitor_name
        paths = set()
        for fname in ['正常数据.xlsx', '异常数据.xlsx']:
            wb = openpyxl.load_workbook(
                r'D:\Desktop\故障定位\data\历史数据\%s' % fname, read_only=True, data_only=True)
            ws = wb['Sheet1']
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            idx = header.index('监测点名称1')
            for row in ws.iter_rows(min_row=2, values_only=True):
                info = parse_monitor_name(row[idx])
                if info and info['line'] == line:
                    paths.add(info['path'])
            wb.close()

        parent = build_tree(paths)
        edges = list(parent.items())  # (child, parent)

        weights = {}
        for child, par in edges:
            weights[child] = edge_weight(par, child)
        total_weight = sum(weights.values())

        code = LINE_CODE[line]
        # 需要清洁编号 - 复用 build_topology 里同样的分配逻辑(DFS前序)
        children = defaultdict(list)
        for node, par in parent.items():
            children[par].append(node)
        for k in children:
            children[k].sort(key=path_sort_key)
        order = []
        stack = list(reversed(children.get(SOURCE, [])))
        while stack:
            cur = stack.pop()
            order.append(cur)
            stack.extend(reversed(children.get(cur, [])))
        clean_id = {node: f'{code}-{i:04d}' for i, node in enumerate(order, 1)}

        for child, par in edges:
            length_km = total_km * weights[child] / total_weight
            is_trunk = len(child) == 1  # 主干/主干补杆,没有小数点分支
            ctype = '干线(正序R1=0.13欧姆/km)' if is_trunk else '支线(正序R1=0.16欧姆/km)'
            p = CONDUCTOR_PARAMS[ctype]
            r_ohm = round(length_km * p['r1_per_km'], 4)
            x_ohm = round(length_km * p['l1_mH_per_km'] * 1e-3 * 2 * 3.141592653589793 * F, 4)
            r0_ohm = round(length_km * ZERO_SEQ['r0_per_km'], 4)
            x0_ohm = round(length_km * ZERO_SEQ['l0_mH_per_km'] * 1e-3 * 2 * 3.141592653589793 * F, 4)
            c1_uF = round(length_km * CAP['c1_uF_per_km'], 6)
            c0_uF = round(length_km * CAP['c0_uF_per_km'], 6)
            out_rows.append({
                'line': line,
                'from_id': 'SOURCE' if par == SOURCE else clean_id[par],
                'to_id': clean_id[child],
                'edge_label': f"{'SOURCE' if par == SOURCE else '.'.join(par)} -> {'.'.join(child)}",
                'length_km_est': round(length_km, 3),
                'conductor_type': ctype,
                'r_ohm': r_ohm,
                'x_ohm': x_ohm,
                'r0_ohm': r0_ohm,
                'x0_ohm': x0_ohm,
                'c1_uF': c1_uF,
                'c0_uF': c0_uF,
            })
        print(f'{line}: 总长度{total_km}km, {len(edges)}条区段, 估算长度合计{sum(r["length_km_est"] for r in out_rows if r["line"]==line):.2f}km')

    out_path = os.path.join(OUT_DIR, 'edges_with_params.csv')
    with open(out_path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=['line', 'from_id', 'to_id', 'edge_label',
                                           'length_km_est', 'conductor_type', 'r_ohm', 'x_ohm',
                                           'r0_ohm', 'x0_ohm', 'c1_uF', 'c0_uF'])
        w.writeheader()
        w.writerows(out_rows)
    print(f'\n已导出: {out_path} ({len(out_rows)}行, 含正序r_ohm/x_ohm + 零序r0_ohm/x0_ohm + 电容c1_uF/c0_uF)')
    print('注意:长度是按编号差加权估算的近似值,不是逐杆塔实测/GIS长度;导线电气参数已换成用户提供的南网10kV文献值。')


if __name__ == '__main__':
    main()
