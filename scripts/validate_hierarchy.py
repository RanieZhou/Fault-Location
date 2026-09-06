# -*- coding: utf-8 -*-
"""用真实故障事件的"同时报警"模式,验证'编号路径前缀=祖先关系'这个假设是否成立。"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from datetime import datetime
from common import parse_monitor_name, path_sort_key, is_ancestor


def to_dt(v):
    if isinstance(v, datetime):
        return v
    return datetime.strptime(str(v).strip(), '%Y-%m-%d %H:%M:%S')

DATA_DIR = r'D:\Desktop\故障定位\data\历史数据'


def load_abnormal_rows():
    wb = openpyxl.load_workbook(os.path.join(DATA_DIR, '异常数据.xlsx'), read_only=True, data_only=True)
    ws = wb['Sheet1']
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(header, r))
        rows.append(d)
    wb.close()
    return rows


def cluster_events(rows, tolerance_sec=15):
    # 按 线路 + 时间接近 聚类成事件
    parsed = []
    for r in rows:
        info = parse_monitor_name(r['监测点名称1'])
        if info is None:
            continue
        info['status'] = r['线路状态']
        info['time'] = to_dt(r['量测时间'])
        parsed.append(info)

    parsed.sort(key=lambda x: (x['line'], x['time']))
    events = []
    cur = []
    for p in parsed:
        if not cur:
            cur = [p]
            continue
        same_line = p['line'] == cur[-1]['line']
        close_time = abs((p['time'] - cur[-1]['time']).total_seconds()) <= tolerance_sec
        if same_line and close_time:
            cur.append(p)
        else:
            events.append(cur)
            cur = [p]
    if cur:
        events.append(cur)
    return events


def main():
    rows = load_abnormal_rows()
    events = cluster_events(rows)
    print(f'共 {len(rows)} 条异常记录, 聚类成 {len(events)} 次故障事件\n')

    n_chain_ok = 0
    n_chain_bad = 0
    for i, ev in enumerate(events, 1):
        line = ev[0]['line']
        t = ev[0]['time']
        pts = sorted({(p['pole'], p['side']) for p in ev})
        paths = sorted({p['path'] for p in ev}, key=path_sort_key)
        # 检查是否构成链:任意两点必须存在祖先关系(相同path也算,通常是大/小两侧)
        is_chain = True
        for a in range(len(paths)):
            for b in range(a + 1, len(paths)):
                pa, pb = paths[a], paths[b]
                if pa == pb:
                    continue
                if not (is_ancestor(pa, pb) or is_ancestor(pb, pa)):
                    is_chain = False
        tag = 'OK  ' if is_chain else 'BAD '
        if is_chain:
            n_chain_ok += 1
        else:
            n_chain_bad += 1
        pts_str = ', '.join(f"{p}({s or '-'})" for p, s in pts)
        print(f'[{i:02d}] {tag}{line} {t}  点位: {pts_str}')

    print(f'\n链式一致: {n_chain_ok} / {len(events)} 个事件, 不一致: {n_chain_bad} 个')


if __name__ == '__main__':
    main()
