# -*- coding: utf-8 -*-
"""
回归校验：用新版 core.fault_locator（含Bug3滑动窗口修复）重新处理原始异常数据，
和旧版 output/fault_location_report.csv 逐条比对，核对事件切分、候选区段、置信度口径。
不改动任何现有输出文件，只打印/写出对比报告供人工判断。
"""
import sys, os, csv
from pathlib import Path
from collections import defaultdict

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import openpyxl
from common import parse_monitor_name
from build_topology import to_dt, IN_SCOPE_LINES, DATA_DIR

from core.fault_locator import group_alarms_into_events, locate_fault
from core.models import FaultLocateRequest


def load_raw_alarms():
    wb = openpyxl.load_workbook(os.path.join(DATA_DIR, '异常数据.xlsx'), read_only=True, data_only=True)
    ws = wb['Sheet1']
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()

    alarms = []
    for r in rows:
        info = parse_monitor_name(r['监测点名称1'])
        if info is None or info['line'] not in IN_SCOPE_LINES:
            continue
        alarms.append({
            'line': info['line'],
            'pole': info['pole'],
            'time': to_dt(r['量测时间']),
            'fault_type': r['线路状态'],
        })
    return alarms


def load_old_report():
    path = ROOT / "output" / "fault_location_report.csv"
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def main():
    alarms = load_raw_alarms()
    events = group_alarms_into_events(alarms)  # 用新滑动窗口分组（Bug3修复版）
    events.sort(key=lambda ev: (ev[0]['line'], ev[0]['time']))

    old_rows = load_old_report()
    old_events = defaultdict(list)
    for r in old_rows:
        old_events[(r['line'], r['time'])].append(r)

    print(f"新引擎分组事件数: {len(events)}")
    print(f"旧报告去重(line,time)事件数: {len(old_events)}  旧报告总行数: {len(old_rows)}\n")

    report_lines = []
    for i, ev in enumerate(events, 1):
        line = ev[0]['line']
        time_str = str(ev[0]['time'])
        alarmed_poles = sorted({p['pole'] for p in ev})
        fault_types = sorted({p['fault_type'] for p in ev})

        req = FaultLocateRequest(line=line, alarm_points=alarmed_poles, fault_type=None)
        result = locate_fault(req)

        new_sections = [(s.from_pole, s.to_pole, s.confidence) for s in result.candidate_sections]
        old = old_events.get((line, time_str), [])

        line_out = f"[{i:02d}] {line}  {time_str}  报警点={alarmed_poles}  故障类型={fault_types}"
        line_out += f"\n     新引擎: frontier={result.frontier_points}  confidence={result.confidence}  sections={new_sections}"
        if old:
            old_frontiers = [r['frontier'] for r in old]
            old_sections = [r['candidate_sections'] for r in old]
            old_notes = [r['note'] for r in old]
            line_out += f"\n     旧报告: frontier={old_frontiers}  sections={old_sections}  note={old_notes}"
        else:
            line_out += "\n     旧报告: 【未找到匹配的(line,time)记录】"
        report_lines.append(line_out)

    out_path = ROOT / "output" / "regression_report.txt"
    out_path.write_text("\n\n".join(report_lines), encoding="utf-8")
    print(f"详细逐条对比已写入: {out_path}")

    # 反向检查：旧报告里有、新分组事件里对不上的 (line,time)
    new_keys = {(ev[0]['line'], str(ev[0]['time'])) for ev in events}
    missing_in_new = [k for k in old_events if k not in new_keys]
    if missing_in_new:
        print(f"\n⚠️ 旧报告中有 {len(missing_in_new)} 个 (line,time) 在新分组里找不到完全匹配:")
        for k in missing_in_new:
            print(f"   {k}  旧行数={len(old_events[k])}")


if __name__ == '__main__':
    main()
