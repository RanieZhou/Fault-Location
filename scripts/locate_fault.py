# -*- coding: utf-8 -*-
"""
确定性的故障区段定位算法(不依赖仿真,直接用故障指示器的"短路"上报逻辑 + 拓扑树)。

核心规则(配电网故障指示器的经典"矩阵法"思路):
  一次单点接地/相间故障发生后,故障电流只会流经"电源 -> 故障点"这一条路径。
  路径上的监测点会报"短路",路径之外(其它分支、故障点下游)的监测点不会。
  所以:把本次事件里所有"短路"上报的点,在拓扑树里找"报警前沿"
  (即:自己报警、但它在树上的直接下游/子节点都没有一起报警的那些点),
  故障就落在"前沿节点 -> 它未报警的下游相邻监测点"这一段区间内。
  如果前沿节点在树上没有下游监测点了(该分支的最后一个装表点),
  只能给出"故障在此点之后,再往下没有监测点可以进一步收窄"的结论——
  这是指示器密度决定的分辨率上限,不是算法的锅。
"""
import sys, os, csv
from collections import defaultdict
from datetime import datetime
import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from common import parse_monitor_name, path_sort_key
from build_topology import build_tree, SOURCE, to_dt, IN_SCOPE_LINES, DATA_DIR, OUT_DIR


def load_all_paths():
    """线路 -> 该线路上所有(历史正常+异常数据里出现过的)监测点path集合。"""
    paths = defaultdict(set)
    for fname in ['正常数据.xlsx', '异常数据.xlsx']:
        wb = openpyxl.load_workbook(os.path.join(DATA_DIR, fname), read_only=True, data_only=True)
        ws = wb['Sheet1']
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = header.index('监测点名称1')
        for row in ws.iter_rows(min_row=2, values_only=True):
            info = parse_monitor_name(row[idx])
            if info and info['line'] in IN_SCOPE_LINES:
                paths[info['line']].add(info['path'])
        wb.close()
    return paths


def children_map(parent):
    ch = defaultdict(list)
    for node, par in parent.items():
        ch[par].append(node)
    for k in ch:
        ch[k].sort(key=path_sort_key)
    return ch


def is_ancestor(a, b, parent):
    """a 是否是 b 在树上的(严格)祖先。"""
    cur = b
    while cur in parent:
        cur = parent[cur]
        if cur == a:
            return True
        if cur == SOURCE:
            return False
    return False


def locate_fault_section(line, alarmed_paths, parent, children):
    """alarmed_paths: 本次事件里报'短路'的 path 集合(同一条线)。
    返回: list of dict, 每个dict是一个'报警前沿'节点推导出的候选故障区段。"""
    results = []
    for node in alarmed_paths:
        has_alarmed_child = any(is_ancestor(node, other, parent) for other in alarmed_paths if other != node)
        if has_alarmed_child:
            continue  # 不是前沿,跳过
        kids = children.get(node, [])
        unalarmed_kids = [k for k in kids if k not in alarmed_paths]
        if not kids:
            results.append({
                'frontier': node, 'candidates': [], 'note': '该分支已无更下游监测点,无法进一步缩小',
            })
        elif len(unalarmed_kids) == 1 and len(kids) == 1:
            results.append({
                'frontier': node, 'candidates': [unalarmed_kids[0]], 'note': '单一候选区段,较有把握',
            })
        else:
            results.append({
                'frontier': node, 'candidates': unalarmed_kids or kids,
                'note': f'该点下游有{len(kids)}条分支,无法仅凭本次数据区分具体哪一条(需要更多信息)',
            })
    return results


def fmt(path):
    return '.'.join(path) if path != SOURCE else 'SOURCE'


def main():
    all_paths = load_all_paths()
    trees = {line: build_tree(all_paths[line]) for line in IN_SCOPE_LINES}
    kids = {line: children_map(trees[line]) for line in IN_SCOPE_LINES}

    wb = openpyxl.load_workbook(os.path.join(DATA_DIR, '异常数据.xlsx'), read_only=True, data_only=True)
    ws = wb['Sheet1']
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    abn_rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()

    parsed = []
    for r in abn_rows:
        info = parse_monitor_name(r['监测点名称1'])
        if info is None or info['line'] not in IN_SCOPE_LINES:
            continue
        info['time'] = to_dt(r['量测时间'])
        info['fault_type'] = r['线路状态']
        parsed.append(info)
    parsed.sort(key=lambda x: (x['line'], x['time']))

    events, cur = [], []
    for p in parsed:
        if cur and p['line'] == cur[-1]['line'] and abs((p['time'] - cur[-1]['time']).total_seconds()) <= 15:
            cur.append(p)
        else:
            if cur:
                events.append(cur)
            cur = [p]
    if cur:
        events.append(cur)

    report_rows = []
    print(f'共 {len(events)} 次历史故障事件,逐一定位:\n')
    for i, ev in enumerate(events, 1):
        line = ev[0]['line']
        alarmed = {p['path'] for p in ev}
        fault_types = sorted({p['fault_type'] for p in ev})
        results = locate_fault_section(line, alarmed, trees[line], kids[line])

        alarmed_str = ', '.join(fmt(p) for p in sorted(alarmed, key=path_sort_key))
        print(f"[{i:02d}] {line}  {ev[0]['time']}  故障类型:{'; '.join(fault_types)}")
        print(f'     报警点: {alarmed_str}')
        for r in results:
            cand_str = ' / '.join(fmt(c) for c in r['candidates']) if r['candidates'] else '(无)'
            print(f"     -> 前沿点 {fmt(r['frontier'])} : 候选故障区段 [{cand_str}]  {r['note']}")
            report_rows.append({
                'event_id': i, 'line': line, 'time': ev[0]['time'], 'fault_types': '; '.join(fault_types),
                'alarmed_points': alarmed_str, 'frontier': fmt(r['frontier']),
                'candidate_sections': cand_str, 'note': r['note'],
            })
        print()

    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, 'fault_location_report.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=['event_id', 'line', 'time', 'fault_types', 'alarmed_points',
                                           'frontier', 'candidate_sections', 'note'])
        w.writeheader()
        w.writerows(report_rows)
    print(f'已导出: {OUT_DIR}\\fault_location_report.csv ({len(report_rows)} 行)')


if __name__ == '__main__':
    main()
